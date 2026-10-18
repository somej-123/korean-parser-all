"""PPTX·XLSX 파서. 시험 문서는 tests/fixtures/make_office.py로 만든다."""

from pathlib import Path

import pytest
from lxml import etree

from korean_parser.ir import Chart, Figure, Heading, PageBreak, Paragraph, Table
from korean_parser.parsers import detect_format, parse
from korean_parser.parsers.xlsx import XlsxParser, format_number

FIX = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="module")
def deck():
    return parse(FIX / "sample.pptx")


@pytest.fixture(scope="module")
def book():
    return parse(FIX / "sample.xlsx")


def cell_texts(t: Table):
    return [[" ".join(p.text for p in c.blocks) for c in r.cells] for r in t.rows]


# -- PPTX ---------------------------------------------------------------------


def test_pptx_detect_and_meta(deck):
    assert detect_format(FIX / "sample.pptx") == "pptx"
    assert deck.meta.page_count == 4 and deck.meta.page_count_source == "layout"
    assert deck.page.width_pt == 720.0
    assert sum(isinstance(b, PageBreak) for b in deck.blocks) == 3


def test_pptx_titles_are_headings_with_slide_page(deck):
    hs = [b for b in deck.blocks if isinstance(b, Heading)]
    assert [h.text for h in hs] == ["2026년 문서 파싱 고도화 계획", "추진 배경", "예산 계획", "분기별 처리량"]
    assert [h.prov.page for h in hs] == [1, 2, 3, 4]
    assert hs[0].prov.bbox is not None  # 자리 표시자 위치는 레이아웃에서 물려받는다


def test_pptx_bullets_with_levels(deck):
    items = [(b.list_info.level, b.text) for b in deck.blocks
             if isinstance(b, Paragraph) and b.list_info and b.prov.page == 2]
    assert items == [(0, "문서 파싱 품질 문제"), (1, "표 구조 손실"), (1, "띄어쓰기 소실"), (0, "처리 속도 문제")]


def test_pptx_table_merges(deck):
    t = next(b for b in deck.blocks if isinstance(b, Table))
    assert cell_texts(t)[:2] == [["구분", "2026년"], ["상반기", "하반기"]]
    assert (t.rows[0].cells[0].row_span, t.rows[0].cells[1].col_span) == (2, 2)


def test_pptx_chart_picture_group_and_notes(deck):
    chart = next(b for b in deck.blocks if isinstance(b, Chart))
    assert chart.categories == ["1분기", "2분기", "3분기"] and chart.series[0].values == [120, 180, 240]
    fig = next(b for b in deck.blocks if isinstance(b, Figure))
    assert fig.alt == "처리 흐름 그림" and fig.image_id in deck.images
    texts = [b.text for b in deck.blocks if isinstance(b, Paragraph)]
    assert "※ 자료: 내부 집계" in texts  # 묶음 안 글상자
    notes = [b for b in deck.blocks if isinstance(b, Paragraph) and b.style_name == "slide_notes"]
    assert [n.text for n in notes] == ["발표 시 띄어쓰기 사례를 함께 보여 줄 것"] and notes[0].prov.page == 2


# -- XLSX ---------------------------------------------------------------------


def test_xlsx_sheets_tables_and_display_values(book):
    assert detect_format(FIX / "sample.xlsx") == "xlsx"
    assert [b.text for b in book.blocks if isinstance(b, Heading)] == ["예산", "일정", "숨김"]
    assert book.warnings == ["숨긴 시트도 넣음: 숨김"]
    t = next(b for b in book.blocks if isinstance(b, Table))
    rows = cell_texts(t)
    assert rows[0] == ["구분", "2026년", "비고"] and rows[1] == ["상반기", "하반기"]
    assert rows[2] == ["인건비", "1,000,000", "1,200,000", "단위: 원"]
    assert rows[3][:3] == ["장비비", "50,000.50", "25%"]
    assert rows[4][:2] == ["착수일", "2026-03-02"]
    assert rows[5][:2] == ["합계", "=B3+B4"]  # 계산 결과가 저장되지 않은 수식은 수식 그대로
    first = t.rows[0].cells
    assert (first[0].row_span, first[1].col_span, first[2].row_span) == (2, 2, 2)
    assert first[0].background == "#DDEBF7" and first[0].blocks[0].inlines[0].style.bold


@pytest.mark.parametrize("value, code, is_date, want", [
    (1234.5, "#,##0", False, "1,235"),
    (0.125, "0.0%", False, "12.5%"),
    (3.0, None, False, "3"),
    (0.1 + 0.2, None, False, "0.3"),
    (46083.0, "yyyy-mm-dd", True, "2026-03-02"),
    (46083.5, "yyyy-mm-dd hh:mm", True, "2026-03-02 12:00"),
    # 단위·통화 글자는 값의 일부다 (예전에는 '1,500,000'만 남았다)
    (1500000, '#,##0"원"', False, "1,500,000원"),
    (1500000, '"₩"#,##0', False, "₩1,500,000"),
    (1500000, "[$₩-412]#,##0", False, "₩1,500,000"),
    (5, "0\개", False, "5개"),
    # 음수·0 부분: 괄호 음수, 회계 형식의 0은 '-'
    (-1500, "#,##0;(#,##0)", False, "(1,500)"),
    (0, '_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-', False, "-"),
    (-1234, '_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-', False, "-1,234"),
    # 분수, 지수, 천 단위 나눔, 글 자리
    (1.5, "# ?/?", False, "1 1/2"),
    (0.3, "?/8", False, "2/8"),
    (12345.678, "0.00E+00", False, "1.23E+04"),
    (1234567, "#,##0,", False, "1,235"),
    (3, "@", False, "3"),
])
def test_number_formats(value, code, is_date, want):
    assert format_number(value, code, is_date) == want


def test_cached_formula_value_wins():
    parser = XlsxParser.__new__(XlsxParser)
    parser.strings, parser.xf, parser.formats, parser.date1904 = [], [], {}, False
    s = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    c = etree.fromstring(f'<c xmlns="{s}" r="B6"><f>B3+B4</f><v>1050000.5</v></c>')
    assert parser._value(c) == "1050000.5"


# -- 자리 표시자 없이 글상자로 만든 PPTX (PptxGenJS 같은 도구) ------------------------


@pytest.fixture(scope="module")
def plain():
    return parse(FIX / "sample_plain.pptx")


def test_plain_slides_repeated_footer_and_page_numbers(plain):
    # 슬라이드마다 같은 자리의 바닥글은 꼬리말로 한 번만, 쪽 번호('3 / 4')는 버린다
    assert [b.text for b in plain.footer] == ["내부 검토용 자료"] and not plain.header
    texts = [b.text for b in plain.blocks if isinstance(b, Paragraph)]
    assert "내부 검토용 자료" not in texts and not any("/ 4" in t for t in texts)
    # 숫자가 슬라이드 번호와 다른 위쪽 이름표('방법 4')는 내용이라 남는다
    assert [t for t in texts if t.startswith("방법")] == ["방법 4", "방법 3", "방법 2", "방법 1"]


def test_plain_slides_title_guess_and_column_order(plain):
    first = [b for b in plain.blocks if isinstance(b, Paragraph) and b.prov.page == 1]
    guessed = [b for b in first if b.heading_guess]
    assert [b.text for b in guessed] == ["검토 항목 1"] and guessed[0].heading_guess.level == 1
    assert "32pt" in guessed[0].heading_guess.reason
    # 두 단: 왼쪽 단을 다 읽고 오른쪽 단으로 (줄마다 번갈아 읽지 않는다)
    order = [b.text for b in first if b.text in ("장점", "+ 처리가 빠르다", "단점", "- 값이 비싸다")]
    assert order == ["장점", "+ 처리가 빠르다", "단점", "- 값이 비싸다"]


def test_plain_slides_path_alt_becomes_file_name(plain):
    fig = next(b for b in plain.blocks if isinstance(b, Figure))
    assert fig.alt == "icon_check.png"


def test_xy_order_reads_cards_one_by_one():
    from korean_parser.parsers.pptx import xy_order

    def item(name, x, y, w=80, h=20):
        return ((1, (x, y, x + w, y + h)), [name])

    # 제목 한 줄 + 카드 세 장(카드마다 위·아래 두 줄). 줄마다가 아니라 카드마다 읽는다
    items = [item("제목", 10, 10, 400), item("1구간", 10, 60), item("2구간", 150, 60), item("3구간", 290, 60),
             item("200만원", 10, 85), item("1,000만원", 150, 85), item("2,500만원", 290, 85)]
    assert [b for _, (b,) in xy_order(items)] == ["제목", "1구간", "200만원", "2구간", "1,000만원", "3구간", "2,500만원"]


def test_xlsx_chart_without_cache_reads_cells(book):
    chart = next(b for b in book.blocks if isinstance(b, Chart))
    assert chart.title == "단계별 진행률" and chart.categories == ["분석", "설계"]
    assert chart.series[0].name == "진행률" and chart.series[0].values == [1.0, 0.45]
    assert chart.prov.page == 2  # '일정' 시트의 표 뒤


# -- 연결선으로 이은 도형 묶음 → 도식 ----------------------------------------------


def test_connected_shapes_become_one_drawing():
    # 실제 화면설계서의 '메뉴구조도'(연결선 18개)가 상자 글만 흩어져 나와 위아래 관계가 사라졌다
    from korean_parser.ir import Drawing
    from korean_parser.render.llm import render_llm

    doc = parse(FIX / "sample_diagram.pptx")
    chart, flow = [b for b in doc.blocks if isinstance(b, Drawing)]
    labels = {s.text_blocks[0].text for s in chart.shapes if s.text_blocks}
    assert labels == {"산업안전", "안전교육", "안전점검", "교육현황"}  # 연결선 없이 쌓은 '교육현황'도 같은 열
    assert sum(s.is_line for s in chart.shapes) == 2 and all(s.id.startswith("s1:") for s in chart.shapes)
    assert min(s.x for s in chart.shapes) == 0 and chart.prov.bbox[0] == pytest.approx(140, abs=1)
    text = render_llm(doc, with_ids=False)
    assert "- 산업안전 — 안전교육" in text and "- 산업안전 — 안전점검" in text
    assert "- 접수 → 처리" in text  # 연결 정보 없이 끝점만 닿은 화살표
    # 도식에 들어간 글은 따로 문단으로 나오지 않고, 이어지지 않은 장식 사각형은 경고만
    assert [b.text for b in doc.blocks if isinstance(b, Paragraph)] == ["메뉴구조도"]
    assert any("장식" in w for w in doc.warnings) and flow.shapes


# -- 원래 배치 화면 (고정 쪽) ------------------------------------------------------


@pytest.fixture(scope="module")
def layout():
    return parse(FIX / "sample_layout.pptx")


def test_fixed_slide_draws_shapes_in_order_with_background(layout):
    # 슬라이드마다 고정 쪽 하나. 마스터의 장식 띠가 맨 아래, 그다음 슬라이드 도형을 원본 그리는 차례대로
    assert len(layout.fixed_pages) == layout.meta.page_count == 1
    fp = layout.fixed_pages[0]
    assert (fp.page, fp.width_pt, fp.height_pt, fp.background) == (1, 720.0, 540.0, "#F5F3EE")
    band, card, box, clear, table = fp.shapes
    assert band.fill == "#1E2761" and band.width == 720 and not band.block_ids  # 마스터 도형은 본문에 넣지 않는다
    assert card.geometry == "roundRect" and card.adjust == {"adj": 5000} and card.fill == "#FBEAEA"
    # 글이 있는 도형은 본문 블록을 id로 가리킨다 (같은 글을 두 번 담지 않는다)
    by_id = {b.id: b for b in layout.blocks}
    assert [by_id[i].text for i in box.block_ids] == ["클라우드 AI", "외부 서버로 전송"]
    assert isinstance(by_id[table.block_ids[0]], Table)
    assert clear.fill is None  # 완전히 투명한(alpha 0) 채우기는 없는 것으로


def test_fixed_slide_text_box_properties(layout):
    box = layout.fixed_pages[0].shapes[2]
    assert box.text_insets_pt == (0, 0, 0, 0) and box.text_wrap is False
    # 크기가 없는 런은 마스터 otherStyle(18pt)과 테마 글꼴을 물려받는다. 본문 블록에는 넣지 않는다
    assert box.text_style.size_pt == 18 and box.text_style.font == "맑은 고딕"
    first, second = (b for b in layout.blocks if isinstance(b, Paragraph) and b.prov.bbox[1] == 120)
    assert first.inlines[0].style is None or first.inlines[0].style.size_pt is None
    # 슬라이드에 적힌 문단 모양(줄 간격, 문단 앞 간격)과 자간은 본문 블록에 그대로
    assert (second.style.line_spacing, second.style.space_before_pt) == (1.2, 6)
    assert second.inlines[0].style.letter_spacing_em == pytest.approx(0.1)


def test_table_style_borders_come_from_file(layout):
    # 칸에 선이 없으면 표 스타일(tableStyles.xml)의 전체 표 선을 쓴다 (실제 화면설계서의 표 43개)
    t = next(b for b in layout.blocks if isinstance(b, Table))
    b = t.rows[0].cells[0].borders
    assert (b.top.color, b.left.color, b.right.color) == ("#9E9E9E",) * 3
    assert t.cell_padding_pt == (3.6, 7.2, 3.6, 7.2)  # PowerPoint 기본 칸 여백


def test_fixed_slide_html_and_flow_unchanged(layout):
    from korean_parser.render.html import render_html

    html = render_html(layout)  # 고정 쪽이 있으면 기본이 원래 배치
    assert 'class="kp-fixed-page kp-fixed-slide" data-page="1"' in html and 'fill="#F5F3EE"' in html
    box = layout.fixed_pages[0].shapes[2]
    assert f'data-kp-id="{box.block_ids[0]}"' in html  # 화면 강조가 본문 블록 id로 이어진다
    assert "kp-nowrap" in html and "padding:0pt 0pt 0pt 0pt" in html
    flow = render_html(layout, layout="flow")
    assert 'class="kp-fixed-page' not in flow and 'data-kp-id="b1"' in flow


def test_fixed_slide_keeps_text_dropped_from_body(plain):
    # 되풀이되어 꼬리말로 한 번만 옮긴 바닥글과 버린 쪽 번호도 원래 배치 화면에는 슬라이드마다 그린다
    def texts(fp):
        by_id = {b.id: b for b in plain.blocks + plain.footer}
        out = []
        for s in fp.shapes:
            out += [getattr(by_id[i], "text", "") for i in s.block_ids] + [b.text for b in s.text_blocks]
        return out

    assert all("내부 검토용 자료" in texts(fp) for fp in plain.fixed_pages)
    assert [t for fp in plain.fixed_pages for t in texts(fp) if "/ 4" in t] == ["1 / 4", "2 / 4", "3 / 4", "4 / 4"]
    # 본문에 없는 글은 id 없이 담는다 (본문 블록 id와 겹치지 않게)
    assert all(not b.id for fp in plain.fixed_pages for s in fp.shapes for b in s.text_blocks)


def test_fixed_slide_draws_diagram_members_in_place():
    # 도식은 한 덩어리가 아니라 구성 도형마다 원래 차례에 그린다 (사이에 놓인 배경 도형이 글을 덮지 않게).
    # 상자 글은 도식 안 문단 id를 가리킨다
    from korean_parser.ir import Drawing, walk_blocks

    doc = parse(FIX / "sample_diagram.pptx")
    drawing_ids = {b.id for b in doc.blocks if isinstance(b, Drawing)}
    shapes = doc.fixed_pages[0].shapes
    assert not any(i in drawing_ids for s in shapes for i in s.block_ids)
    inner = {b.id: b for b in walk_blocks(doc.blocks)}
    labels = [inner[i].text for s in shapes for i in s.block_ids]
    assert {"산업안전", "안전교육", "교육현황", "접수", "처리"} <= set(labels)
    assert sum(s.is_line for s in shapes) == 3 and any(s.tail_arrow for s in shapes)


def _xlsx(tmp_path, sheet_xml: str, parts: dict[str, str], sheet_rels: str = "") -> Path:
    """시트 하나짜리 최소 XLSX (시트 XML, 더할 부분들, 시트 관계)."""
    import zipfile

    ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" ' \
         'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    path = tmp_path / "t.xlsx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="xml" ContentType="application/xml"/><Default Extension="rels" '
                   'ContentType="application/vnd.openxmlformats-package.relationships+xml"/></Types>')
        z.writestr("_rels/.rels", f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   f'<Relationship Id="rId1" Type="{rel}/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        z.writestr("xl/workbook.xml", f'<workbook {ns}><sheets><sheet name="예산" sheetId="1" r:id="rId1"/></sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   f'<Relationship Id="rId1" Type="{rel}/worksheet" Target="worksheets/sheet1.xml"/></Relationships>')
        z.writestr("xl/worksheets/sheet1.xml", f"<worksheet {ns}>{sheet_xml}</worksheet>")
        if sheet_rels:
            z.writestr("xl/worksheets/_rels/sheet1.xml.rels",
                       f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{sheet_rels}</Relationships>')
        for name, xml in parts.items():
            z.writestr(name, xml)
    return path


def test_xlsx_comments_links_header_footer_shapes(tmp_path):
    # 예전에는 칸 메모, 하이퍼링크 주소, 인쇄 머리글·바닥글, 글상자 글이 경고 없이 빠졌다
    from korean_parser.ir import Drawing, NoteRef, walk_blocks

    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    sheet = ('<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>누리집</t></is></c>'
             '<c r="B1" t="inlineStr"><is><t>본문표</t></is></c></row>'
             '<row r="2"><c r="A2"><v>3</v></c></row></sheetData>'
             '<hyperlinks><hyperlink ref="A1" r:id="rId9"/><hyperlink ref="B1" location="\'요약\'!A1"/></hyperlinks>'
             '<headerFooter><oddHeader>&amp;L&amp;"맑은 고딕,굵게"&amp;12사업 보고서&amp;R&amp;A</oddHeader>'
             '<oddFooter>&amp;C&amp;P / &amp;N쪽</oddFooter></headerFooter>'
             '<drawing r:id="rId2"/><oleObjects><oleObject progId="Word.Document.12" shapeId="1"/></oleObjects>')
    rels = (f'<Relationship Id="rId9" Type="{rel}/hyperlink" Target="https://www.data.go.kr" TargetMode="External"/>'
            f'<Relationship Id="rId1" Type="{rel}/comments" Target="../comments1.xml"/>'
            f'<Relationship Id="rId2" Type="{rel}/drawing" Target="../drawings/drawing1.xml"/>')
    comments = ('<comments xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><authors><author>a</author>'
                '</authors><commentList><comment ref="B1" authorId="0"><text><r><t>단위: 천원</t></r></text></comment>'
                '<comment ref="C5" authorId="0"><text><t>빈 칸 메모</t></text></comment></commentList></comments>')
    a = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    drawing = ('<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" ' + a + '>'
               '<xdr:twoCellAnchor><xdr:from/><xdr:to/><xdr:grpSp><xdr:sp><xdr:spPr><a:xfrm><a:off x="0" y="0"/>'
               '<a:ext cx="1270000" cy="381000"/></a:xfrm><a:prstGeom prst="rect"/></xdr:spPr><xdr:txBody><a:p><a:r>'
               '<a:t>※ 잠정치</a:t></a:r></a:p></xdr:txBody></xdr:sp></xdr:grpSp></xdr:twoCellAnchor></xdr:wsDr>')
    path = _xlsx(tmp_path, sheet, {"xl/comments1.xml": comments, "xl/drawings/drawing1.xml": drawing}, rels)
    doc = parse(path)
    table = next(b for b in doc.blocks if isinstance(b, Table))
    a1, b1 = table.rows[0].cells[0].blocks[0], table.rows[0].cells[1].blocks[0]
    assert a1.inlines[0].link == "https://www.data.go.kr" and b1.inlines[0].link == "#'요약'!A1"
    assert isinstance(b1.inlines[-1], NoteRef)
    assert [n.blocks[0].text for n in doc.notes.values()] == ["단위: 천원", "빈 칸 메모"]
    assert all(n.kind == "comment" for n in doc.notes.values())
    assert len(table.rows) == 5 and len(table.rows[0].cells) == 3  # 메모만 있는 C5까지 표에 남는다
    assert [p.text for p in doc.header] == ["사업 보고서 예산"] and [p.text for p in doc.footer] == ["/ 쪽"]
    drawing_block = next(b for b in doc.blocks if isinstance(b, Drawing))
    assert [p.text for p in walk_blocks([drawing_block]) if isinstance(p, Paragraph)] == ["※ 잠정치"]
    assert any("OLE" in w for w in doc.warnings)


def test_xlsx_threaded_comments_win_and_no_cell_limit(tmp_path, monkeypatch):
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rows = "".join(f'<row r="{r}"><c r="A{r}"><v>{r}</v></c></row>' for r in range(1, 31))
    tc = ('<ThreadedComments xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments">'
          '<threadedComment ref="A1" id="{1}" personId="{9}"><text>확인 필요</text></threadedComment>'
          '<threadedComment ref="A1" id="{2}" personId="{9}" parentId="{1}"><text>확인함</text></threadedComment>'
          '</ThreadedComments>')
    legacy = ('<comments xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><commentList>'
              '<comment ref="A1"><text><t>[스레드 메모] 이 버전의 Excel에서는…</t></text></comment></commentList></comments>')
    rels = (f'<Relationship Id="rId1" Type="{rel}/comments" Target="../comments1.xml"/>'
            f'<Relationship Id="rId2" Type="http://schemas.microsoft.com/office/2017/10/relationships/threadedComment" '
            f'Target="../threadedComments/threadedComment1.xml"/>')
    path = _xlsx(tmp_path, f"<sheetData>{rows}</sheetData>",
                 {"xl/comments1.xml": legacy, "xl/threadedComments/threadedComment1.xml": tc}, rels)
    doc = parse(path)
    assert [p.text for n in doc.notes.values() for p in n.blocks] == ["확인 필요", "확인함"]
    assert len(next(b for b in doc.blocks if isinstance(b, Table)).rows) == 30 and not doc.warnings
    monkeypatch.setenv("KP_XLSX_MAX_CELLS", "10")  # 메모리가 모자란 환경: 상한을 두면 앞부분만 넣고 경고
    doc = parse(path)
    assert len(next(b for b in doc.blocks if isinstance(b, Table)).rows) == 10
    assert any("KP_XLSX_MAX_CELLS" in w for w in doc.warnings)


def test_pptx_comments_hidden_slide_link_math_ole(tmp_path):
    # PPTX 점검(v1.4.0)에서 찾은 것: 슬라이드 메모(옛·새 형식, 답글), 하이퍼링크 주소, 수식(a14:m), OLE 개체의
    # 대신 그림이 모두 경고 없이 빠지거나 경고만 남았다. 숨긴 슬라이드는 읽되 알린다
    import base64
    import io
    import re
    import zipfile

    from pptx import Presentation
    from pptx.util import Pt

    from korean_parser.ir import Equation, Figure, Text

    prs = Presentation()
    for i in range(2):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f"슬라이드 {i + 1}"
        box = s.shapes.add_textbox(Pt(50), Pt(150), Pt(400), Pt(100)).text_frame
        box.text = "누리집 안내"
        box.paragraphs[0].runs[0].hyperlink.address = "https://example.go.kr/"
    buf = io.BytesIO()
    prs.save(buf)
    out = tmp_path / "audit.pptx"
    png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGP4DwABAQEABRjYTgAAAABJRU5ErkJggg==")
    ns = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
          'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')
    math = ('<p:sp><p:nvSpPr><p:cNvPr id="90" name="수식"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm>'
            '<a:off x="635000" y="3810000"/><a:ext cx="2540000" cy="635000"/></a:xfrm></p:spPr><p:txBody><a:bodyPr/><a:p>'
            '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
            '<mc:Choice xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main" Requires="a14"><a14:m>'
            '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath><m:f><m:num><m:r>'
            '<m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f></m:oMath></m:oMathPara></a14:m>'
            '</mc:Choice><mc:Fallback><a:r><a:t> </a:t></a:r></mc:Fallback></mc:AlternateContent></a:p></p:txBody></p:sp>')
    ole = ('<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="91" name="개체"/><p:cNvGraphicFramePr/><p:nvPr/>'
           '</p:nvGraphicFramePr><p:xfrm><a:off x="635000" y="4445000"/><a:ext cx="2540000" cy="1270000"/></p:xfrm>'
           '<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/presentationml/2006/ole">'
           '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
           '<mc:Choice Requires="v"><p:oleObj r:id="rIdOle" progId="Excel.Sheet.12"><p:embed/></p:oleObj></mc:Choice>'
           '<mc:Fallback><p:oleObj r:id="rIdOle" progId="Excel.Sheet.12"><p:embed/><p:pic><p:nvPicPr><p:cNvPr id="2" '
           'name="개체 그림"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr><p:blipFill><a:blip r:embed="rIdOleImg"/></p:blipFill>'
           '<p:spPr/></p:pic></p:oleObj></mc:Fallback></mc:AlternateContent></a:graphicData></a:graphic></p:graphicFrame>')
    with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/slides/slide2.xml":
                xml = re.sub(r"<p:sld\b", '<p:sld show="0"', data.decode("utf-8"), count=1)
                data = xml.replace("</p:spTree>", math + ole + "</p:spTree>").encode("utf-8")
            elif item.filename == "ppt/slides/_rels/slide1.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rIdCm" Type="http://schemas.openxmlformats.org'
                                    b'/officeDocument/2006/relationships/comments" Target="../comments/comment1.xml"/>'
                                    b"</Relationships>")
            elif item.filename == "ppt/slides/_rels/slide2.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rIdCm2" Type="http://schemas.microsoft.com/'
                                    b'office/2018/10/relationships/comments" Target="../comments/modernComment_1.xml"/>'
                                    b'<Relationship Id="rIdOle" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                                    b'relationships/package" Target="../embeddings/sheet1.xlsx"/>'
                                    b'<Relationship Id="rIdOleImg" Type="http://schemas.openxmlformats.org/officeDocument/'
                                    b'2006/relationships/image" Target="../media/ole1.png"/></Relationships>')
            zout.writestr(item, data)
        zout.writestr("ppt/comments/comment1.xml",
                      f'<p:cmLst {ns}><p:cm authorId="0" idx="1"><p:pos x="10" y="10"/><p:text>옛 메모 글</p:text></p:cm>'
                      "</p:cmLst>")
        body = '<p188:txBody><a:bodyPr/><a:p><a:r><a:t>{}</a:t></a:r></a:p></p188:txBody>'
        zout.writestr("ppt/comments/modernComment_1.xml",
                      f'<p188:cmLst {ns} xmlns:p188="http://schemas.microsoft.com/office/powerpoint/2018/8/main">'
                      '<p188:cm id="{1}" authorId="{2}"><p188:replyLst><p188:reply id="{3}" authorId="{2}">'
                      + body.format("답글") + "</p188:reply></p188:replyLst>" + body.format("새 메모 글")
                      + "</p188:cm></p188:cmLst>")
        zout.writestr("ppt/embeddings/sheet1.xlsx", b"")
        zout.writestr("ppt/media/ole1.png", png)
    doc = parse(out)
    texts = [i for b in doc.iter_blocks() for i in getattr(b, "inlines", []) if isinstance(i, Text)]
    assert any(t.text == "누리집 안내" and t.link == "https://example.go.kr/" for t in texts)
    memos = [[p.text for p in n.blocks] for n in doc.notes.values() if n.kind == "comment"]
    assert memos == [["옛 메모 글"], ["새 메모 글", "답글"]]
    assert [n.blocks[0].prov.page for n in doc.notes.values()] == [1, 2]
    eqs = [i for b in doc.iter_blocks() for i in getattr(b, "inlines", []) if isinstance(i, Equation)]
    assert eqs and "frac" in eqs[0].script
    figs = [b for b in doc.iter_blocks() if isinstance(b, Figure)]
    assert figs and figs[0].prov.page == 2
    assert "숨긴 슬라이드도 넣음 (2쪽)" in doc.warnings
    assert "끼워 넣은 OLE 개체(다른 파일)는 그림으로만 넣음 (PPTX)" in doc.warnings
