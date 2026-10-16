"""목차가 있는 시험 문서(toc.docx, toc.pdf) 생성.

한국 제안요청서의 목차처럼 '번호 제목 ······ 쪽'을 오른쪽 탭(가운뎃점 채움)으로 쓴 줄을 늘어놓는다.
Docling은 이런 영역을 목차(document_index) 표로 보는데, 표 구조 모델에 맡기면 번호와 제목이 다른 칸으로 갈라졌다.

    python tests/fixtures/make_toc.py   # → tests/fixtures/toc.docx, toc.pdf (LibreOffice 필요)
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Pt

HERE = Path(__file__).parent
ENTRIES = [("I. 사업 개요", 1), ("1. 사업 개요", 1), ("2. 사업목표 및 필요성", 1), ("3. 사업 범위", 2),
           ("II. 제안요청 내용", 3), ("1. 사업 내용", 3), ("2. 사업 추진 체계", 9), ("3. 사업 보고 및 성과물 제출", 10),
           ("III. 제안 일반사항", 12), ("1. 입찰참가자격", 12), ("2. 제안서 제출 및 접수", 13), ("IV. 제안서 작성 안내", 21)]


def main() -> None:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(12)
    run = doc.add_paragraph().add_run("< 목 차 >")
    run.bold, run.font.size = True, Pt(16)
    for text, page in ENTRIES:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Pt(450), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.MIDDLE_DOT)
        if not text.startswith(("I", "V")):
            p.paragraph_format.left_indent = Pt(12)
        p.add_run(f"{text}\t{page}")
    doc.add_page_break()
    doc.add_paragraph().add_run("I. 사업 개요").bold = True
    doc.add_paragraph("본 사업은 기관 내부 문서를 통합 관리하고 검색 품질을 높이기 위하여 추진한다.")
    src = HERE / "toc.docx"
    doc.save(src)
    soffice = shutil.which("soffice") or "C:/Program Files/LibreOffice/program/soffice.exe"
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run([soffice, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--convert-to", "pdf",
                        "--outdir", str(HERE), str(src)], check=True, capture_output=True)


if __name__ == "__main__":
    main()
