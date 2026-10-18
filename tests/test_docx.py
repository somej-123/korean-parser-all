"""DOCX → IR → HTML/LLM 파이프라인 테스트 (tests/fixtures/korean_report.docx 기준)."""

from pathlib import Path

import pytest

from korean_parser.ir import Chart, Document, Drawing, Heading, ImageRef, Paragraph, Table
from korean_parser.parsers import DocumentError, parse
from korean_parser.render.html import render_html
from korean_parser.render.llm import render_llm

FIXTURE = Path(__file__).parent / "fixtures" / "korean_report.docx"


@pytest.fixture(scope="module")
def doc() -> Document:
    return parse(FIXTURE)


def blocks_of(doc, kind):
    return [b for b in doc.iter_blocks() if isinstance(b, kind)]


def test_metadata_and_page(doc):
    assert doc.meta.title == "2026년 문서 파싱 고도화 계획"
    assert doc.meta.source_format == "docx"
    assert doc.page.width_pt > 500
    assert doc.header and "대외비" in doc.header[0].text


def test_headings(doc):
    headings = [(h.level, h.text) for h in blocks_of(doc, Heading)]
    assert headings[0] == (1, "2026년 문서 파싱 고도화 계획")
    assert (1, "3. 예산 계획") in headings


def test_korean_numbering_markers(doc):
    markers = [(p.list_info.level, p.list_info.marker, p.text) for p in blocks_of(doc, Paragraph) if p.list_info]
    assert markers[:5] == [
        (0, "1.", "통합 문서 모델 구축"),
        (1, "가.", "DOCX 파서 개발"),
        (2, "1)", "문단과 서식"),
        (3, "가)", "자동 번호"),
        (3, "나)", "셀 병합"),
    ]
    assert (1, "나.", "HWPX 파서 개발") in markers
    assert (0, "2.", "화면 표시 품질 검증") in markers


def test_outline_from_manual_bullets(doc):
    outline = [(p.outline.level, p.outline.marker) for p in blocks_of(doc, Paragraph) if p.outline]
    assert outline[:4] == [(0, "□"), (1, "○"), (2, "-"), (2, "-")]
    assert (2, "※") in outline


def test_merged_table(doc):
    table = blocks_of(doc, Table)[0]
    first = table.rows[0].cells
    assert (first[0].row_span, first[0].col_span) == (2, 1)
    assert (first[1].row_span, first[1].col_span) == (1, 3)
    assert first[0].background == "#D9E2F3"
    assert len(table.rows) == 5 and len(table.rows[1].cells) == 3  # 세로 병합된 칸은 빠진다
    assert first[0].borders and first[0].borders.top


def test_image_and_hyperlink(doc):
    images = [i for p in blocks_of(doc, Paragraph) for i in p.inlines if isinstance(i, ImageRef)]
    assert len(images) == 1 and doc.images[images[0].image_id].mime == "image/png"
    links = [i.link for p in blocks_of(doc, Paragraph) for i in p.inlines if getattr(i, "link", None)]
    assert "https://github.com/docling-project/docling" in links


def test_flowchart_shapes_and_connections(doc):
    drawing = blocks_of(doc, Drawing)[0]
    boxes = {s.id: s for s in drawing.shapes if not s.is_line}
    lines = [s for s in drawing.shapes if s.is_line]
    assert {s.geometry for s in boxes.values()} >= {"flowChartTerminator", "flowChartProcess", "flowChartDecision"}
    assert len(lines) == 4 and all(ln.tail_arrow for ln in lines)
    assert ("11", "12") in {(ln.connect_from, ln.connect_to) for ln in lines}
    texts = {" ".join(b.text for b in s.text_blocks) for s in boxes.values()}
    assert {"신청", "부서 검토", "적합?", "승인", "반려"} <= texts


def test_html_rendering(doc):
    html = render_html(doc)
    assert html.startswith("<!doctype html>")
    assert 'rowspan="2"' in html and 'colspan="3"' in html
    assert "<svg" in html and "foreignObject" in html
    assert "data:image/png;base64," in html
    assert 'class="kp-marker"' in html and ">가.</span>" in html
    assert "<script" not in html.lower()


def test_llm_rendering(doc):
    text = render_llm(doc)
    assert "# 2026년 문서 파싱 고도화 계획" in text
    assert '<td rowspan="2">구분</td><td colspan="3">2026년</td>' in text  # 병합 표는 HTML
    assert "  가. DOCX 파서 개발" in text
    assert "- 신청 → 부서 검토" in text and "- 적합? → 반려" in text
    assert "<!-- b" in text  # 블록 id


def test_json_roundtrip(doc):
    again = Document.model_validate_json(doc.model_dump_json())
    assert render_llm(again) == render_llm(doc)


def test_html_escapes_text_and_blocks_unsafe_links():
    from korean_parser.ir import Text

    d = Document(blocks=[Paragraph(id="b1", inlines=[
        Text(text="<script>alert(1)</script>"),
        Text(text="링크", link="javascript:alert(1)"),
    ])])
    html = render_html(d)
    assert "<script>alert" not in html and "&lt;script&gt;" in html
    assert "javascript:" not in html


def test_errors_are_readable(tmp_path):
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"PK\x03\x04 not really a zip")
    with pytest.raises(DocumentError, match="손상된 파일"):
        parse(broken)
    ole = tmp_path / "encrypted.docx"
    ole.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100)
    with pytest.raises(DocumentError, match="암호화"):
        parse(ole)


def test_chartex_data():
    # Office 2016 신형 차트(cx:chartSpace)는 예전에 '지원하지 않는 그래픽 요소: chartex' 경고만 남겼다.
    # 데이터 묶음(분류·값)과 계열(이름, 데이터 번호, 종류)을 읽는다. 여러 단계 분류는 위 단계부터 ' / '로 잇는다
    from lxml import etree

    from korean_parser.parsers.docx.drawing import read_chart

    cx = "http://schemas.microsoft.com/office/drawing/2014/chartex"
    xml = f"""<cx:chartSpace xmlns:cx="{cx}"><cx:chartData><cx:data id="0">
      <cx:strDim type="cat"><cx:lvl ptCount="2"><cx:pt idx="0">잎1</cx:pt><cx:pt idx="1">잎2</cx:pt></cx:lvl>
        <cx:lvl ptCount="2"><cx:pt idx="0">줄기</cx:pt><cx:pt idx="1">줄기</cx:pt></cx:lvl></cx:strDim>
      <cx:numDim type="size"><cx:lvl ptCount="2"><cx:pt idx="0">22</cx:pt><cx:pt idx="1">12.5</cx:pt></cx:lvl></cx:numDim>
    </cx:data></cx:chartData><cx:chart><cx:title><cx:tx><cx:txData><cx:v>예산 구성</cx:v></cx:txData></cx:tx></cx:title>
    <cx:plotArea><cx:plotAreaRegion><cx:series layoutId="sunburst"><cx:tx><cx:txData><cx:v>금액</cx:v></cx:txData></cx:tx>
    <cx:dataId val="0"/></cx:series></cx:plotAreaRegion></cx:plotArea></cx:chart></cx:chartSpace>"""
    chart = read_chart(etree.fromstring(xml), 300, 200)
    assert (chart.chart_type, chart.title) == ("sunburst", "예산 구성")
    assert chart.categories == ["줄기 / 잎1", "줄기 / 잎2"]
    assert [(s.name, s.values) for s in chart.series] == [("금액", [22.0, 12.5])]


def test_omml_equation_to_latex():
    # Word 수식(OMML)은 예전에 통째로 빠졌다. 구조 요소를 LaTeX 문자열로 옮겨 Equation에 둔다
    from lxml import etree

    from korean_parser.parsers.docx.omml import omml_to_latex

    m = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'

    def r(t):
        return f"<m:r><m:t>{t}</m:t></m:r>"

    xml = (f"<m:oMath {m}>{r('x')}{r('=')}<m:f><m:num>{r('-b')}<m:rad><m:radPr><m:degHide m:val='1'/></m:radPr>"
           f"<m:deg/><m:e>{r('D')}</m:e></m:rad></m:num><m:den>{r('2a')}</m:den></m:f>"
           f"<m:nary><m:naryPr><m:chr m:val='∑'/></m:naryPr><m:sub>{r('i=1')}</m:sub><m:sup>{r('n')}</m:sup>"
           f"<m:e><m:sSub><m:e>{r('x')}</m:e><m:sub>{r('i')}</m:sub></m:sSub></m:e></m:nary></m:oMath>")
    assert omml_to_latex(etree.fromstring(xml)) == r"x=\frac{-b\sqrt{D}}{2a}\sum_{i=1}^{n}{{x}_{i}}"


def test_symbol_font_bullets():
    # LibreOffice가 DOCX로 내보낸 □ ○ 글머리표는 MT Extra 0x57, Wingdings 2 0x81이다 (LibreOffice fontcvt 표)
    from korean_parser.parsers.docx.symbols import map_symbol

    assert map_symbol("MT Extra", "\uf057") == "□" and map_symbol("Wingdings 2", "\uf081") == "○"
    assert map_symbol("Wingdings", "\uf06f") == "□" and map_symbol("Wingdings 3", "\uf075") == "▶"


def _docx(tmp_path, body: str, headers: dict[str, str]) -> Path:
    """최소 DOCX를 만든다 (본문 XML, 머리말 이름 → 머리말 XML)."""
    import zipfile

    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    ns = (f'{w} xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
          'xmlns:v="urn:schemas-microsoft-com:vml"')
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    path = tmp_path / "t.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="xml" ContentType="application/xml"/><Default Extension="rels" '
                   'ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/word/document.xml" '
                   'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr("_rels/.rels", f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   f'<Relationship Id="rId1" Type="{rel}/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr("word/_rels/document.xml.rels",
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   + "".join(f'<Relationship Id="{n}" Type="{rel}/header" Target="{n}.xml"/>' for n in headers)
                   + "</Relationships>")
        z.writestr("word/document.xml", f"<w:document {ns}><w:body>{body}</w:body></w:document>")
        for n, xml in headers.items():
            z.writestr(f"word/{n}.xml", f"<w:hdr {ns}>{xml}</w:hdr>")
    return path


def _p(text: str) -> str:
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def test_vml_group_reads_every_text_box(tmp_path):
    # PDF를 바꾼 DOCX는 한 w:pict의 v:group에 글상자가 여럿 있다. 예전에는 첫 글상자만 읽어 나머지 글이 경고 없이 빠졌다
    box = '<v:shape style="position:absolute;left:{x};top:0;width:400;height:100"><v:textbox><w:txbxContent>{p}' \
          '</w:txbxContent></v:textbox></v:shape>'
    group = ('<v:group style="width:100pt;height:20pt" coordsize="1000,200">'
             + box.format(x=0, p=_p("첫 상자")) + box.format(x=500, p=_p("둘째 상자")) + "</v:group>")
    path = _docx(tmp_path, f"<w:p><w:r><w:pict>{group}</w:pict></w:r></w:p><w:sectPr/>", {})
    drawing = next(b for b in parse(path).blocks if isinstance(b, Drawing))
    texts = [s.text_blocks[0].text for s in drawing.shapes]
    assert texts == ["첫 상자", "둘째 상자"]
    # 묶음 좌표(1000 × 200)를 묶음 크기(100pt × 20pt)로 바꾼다
    assert [(s.x, s.width) for s in drawing.shapes] == [(0, 40), (50, 40)]


def test_other_section_headers_are_kept(tmp_path):
    # 화면용 머리말은 마지막 구역의 기본 머리말이다. 다른 구역·짝수 쪽 머리말도 글이 다르면 뒤에 담는다
    sect1 = '<w:sectPr><w:headerReference w:type="default" r:id="h1"/><w:headerReference w:type="even" r:id="h2"/></w:sectPr>'
    body = f"<w:p><w:pPr>{sect1}</w:pPr><w:r><w:t>본문</w:t></w:r></w:p>" \
           '<w:sectPr><w:headerReference w:type="default" r:id="h3"/></w:sectPr>'
    path = _docx(tmp_path, body, {"h1": _p("제1편"), "h2": _p("제2편"), "h3": _p("제3편"), "h4": _p("제1편")})
    assert [b.text for b in parse(path).header] == ["제3편", "제1편", "제2편"]


def test_format_comes_from_content_not_extension(tmp_path):
    # 공공데이터포털에 '.hwpx'로 올라간 PDF를 HWPX로 열다가 '손상된 파일'로 거절했다. 내용으로 형식을 정하고 경고를 남긴다
    import shutil

    from korean_parser.parsers import detect_format

    src = Path(__file__).parent / "fixtures" / "korean_report.docx"
    wrong = tmp_path / "report.hwpx"
    shutil.copy(src, wrong)
    assert detect_format(wrong) == "docx"
    doc = parse(wrong)
    assert doc.meta.source_format == "docx"
    assert doc.warnings[0] == "확장자는 .hwpx이지만 내용은 DOCX라서 DOCX로 읽음"
    fake_pdf = tmp_path / "notice.hwpx"
    fake_pdf.write_bytes(b"\n%PDF-1.6\n%\xe2\xe3\xcf\xd3\n")
    assert detect_format(fake_pdf) == "pdf"
    # 내용으로 알 수 없으면 확장자를 따른다 (그다음 그 형식의 파서가 손상 여부를 알린다)
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a zip")
    assert detect_format(broken) == "docx"


def test_word_comments_and_replies(tmp_path):
    # Word 메모는 예전에 '넣지 않음' 경고만 남겼다. 메모 참조 자리에 NoteRef, 답글은 부모 메모 안에 붙인다
    import zipfile

    from korean_parser.ir import NoteRef

    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" ' \
        'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"'
    body = ('<w:p><w:commentRangeStart w:id="0"/><w:r><w:t>사업 예산은 12억 원이다.</w:t></w:r><w:commentRangeEnd w:id="0"/>'
            '<w:r><w:commentReference w:id="0"/></w:r></w:p><w:sectPr/>')
    path = _docx(tmp_path, body, {})
    comments = (f'<w:comments {w}><w:comment w:id="0" w:author="검토자"><w:p w14:paraId="0000000A"><w:r>'
                '<w:t>부가세 포함인지 확인</w:t></w:r></w:p></w:comment>'
                '<w:comment w:id="1" w:author="작성자"><w:p w14:paraId="0000000B"><w:r><w:t>포함입니다</w:t></w:r></w:p>'
                '</w:comment></w:comments>')
    ext = ('<w15:commentsEx xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml">'
           '<w15:commentEx w15:paraId="0000000A"/><w15:commentEx w15:paraId="0000000B" w15:paraIdParent="0000000A"/>'
           '</w15:commentsEx>')
    with zipfile.ZipFile(path, "a") as z:
        z.writestr("word/comments.xml", comments)
        z.writestr("word/commentsExtended.xml", ext)
    doc = parse(path)
    refs = [i for i in doc.blocks[0].inlines if isinstance(i, NoteRef)]
    assert [r.label for r in refs] == ["메모1"]
    note = doc.notes[refs[0].note_id]
    assert note.kind == "comment" and [b.text for b in note.blocks] == ["부가세 포함인지 확인", "포함입니다"]
    assert not any("메모" in w for w in doc.warnings)


def test_docx_moved_text_ruby_hidden_altchunk_ole(tmp_path):
    # DOCX 점검에서 찾은 빈틈: 변경 추적의 옮긴 글(w:moveTo)과 덧말(w:ruby)의 본말이 빠졌고, 숨김 서식 글과
    # 끼워 넣은 내용(w:altChunk), OLE 개체는 아무 표시 없이 빠졌다
    import zipfile

    body = ('<w:p><w:moveFrom w:id="1"><w:r><w:t>옮기기 전</w:t></w:r></w:moveFrom>'
            '<w:moveTo w:id="2"><w:r><w:t>옮긴 글</w:t></w:r></w:moveTo>'
            '<w:r><w:rPr><w:vanish/></w:rPr><w:t>숨긴 글</w:t></w:r>'
            '<w:ruby><w:rt><w:r><w:t>한자</w:t></w:r></w:rt><w:rubyBase><w:r><w:t>漢字</w:t></w:r></w:rubyBase></w:ruby></w:p>'
            '<w:altChunk r:id="rIdA"/><w:p><w:r><w:object><v:shape><o:OLEObject '
            'xmlns:o="urn:schemas-microsoft-com:office:office" ProgID="Excel.Sheet.12"/></v:shape></w:object></w:r></w:p>'
            "<w:sectPr/>")
    path = _docx(tmp_path, body, {})
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    files["word/_rels/document.xml.rels"] = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rIdA" Type="{rel}/aFChunk" Target="chunk.html"/></Relationships>').encode()
    files["word/chunk.html"] = "<html><body><p>끼워 넣은 문단</p><script>x()</script></body></html>".encode("utf-8")
    with zipfile.ZipFile(path, "w") as z:
        for n, data in files.items():
            z.writestr(n, data)
    doc = parse(path)
    texts = [b.text for b in doc.blocks if isinstance(b, Paragraph)]
    assert texts[0] == "옮긴 글漢字(한자)" and "끼워 넣은 문단" in texts
    assert "숨김 서식 글 4자는 넣지 않음" in doc.warnings
    assert any("OLE" in w for w in doc.warnings)


def test_big_type_only_title_and_placeholder_heading(tmp_path):
    # 실제 양식 DOCX: '개정 이력'은 번호·스타일 없이 20pt 굵게만 썼다. 크기가 적힌 글이 이 줄뿐이라 본문 크기가 20pt로
    # 잡혀 못 찾았다 (크기 없는 글은 문서 기본 크기로 센다). '…'뿐인 제목 스타일 문단은 제목이 아니다
    from docx import Document as Docx
    from docx.shared import Pt

    from korean_parser.ir import Heading, Paragraph

    d = Docx()
    run = d.add_paragraph().add_run("개정 이력")
    run.bold, run.font.size = True, Pt(20)
    d.add_paragraph("본문 글입니다. " * 5)
    d.add_heading("…", level=1)
    big_number = d.add_paragraph().add_run("2026")
    big_number.font.size = Pt(24)
    path = tmp_path / "big.docx"
    d.save(path)
    doc = parse(path)
    first = doc.blocks[0]
    assert isinstance(first, Paragraph) and first.heading_guess and first.heading_guess.level == 1
    assert "큰 글씨만" in first.heading_guess.reason
    assert not any(isinstance(b, Heading) for b in doc.blocks)  # '…' 제목 스타일 문단은 문단으로
    assert not any(b.heading_guess for b in doc.blocks[1:] if isinstance(b, Paragraph))  # 숫자뿐인 큰 글씨는 아니다
